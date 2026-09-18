# =============================================================================
# DOMINIUM | MAPA DE RESPONSABILIDADE
#
# IMPERIUM
# - SIM - exportacao em planilha (.xlsx) de baixas efetuadas no Imperium.
#
# TOA
# - NAO - este arquivo nao consulta nem automatiza o TOA.
#
# DOMINIUM COMPARTILHADO
# - Apoio a relatorios e formatacao operacional.
#
# Categoria deste arquivo: IMPERIUM.
# Mapa completo: MAPA_DOMINIUM_IMPERIUM_TOA.md
# =============================================================================
from __future__ import annotations

import datetime as dt
import io
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def format_materials_readable(materials: Any) -> str:
    """Converte lista ou dados de materiais em uma string legivel.

    Exemplo: "100m CABO 2FO, 6m FITA ISOLANTE, 2x CONECTOR REUTILIZAVEL"
    """
    if not materials:
        return "-"
    if isinstance(materials, str):
        return materials.strip() or "-"
    if not isinstance(materials, (list, tuple)):
        return "-"

    formatted_items = []
    for item in materials:
        if not isinstance(item, dict):
            continue
        qty = item.get("quantity") if item.get("quantity") is not None else item.get("qty")
        if qty is None:
            qty = item.get("stock_quantity") or 0

        try:
            qty_num = float(qty)
            if qty_num.is_integer():
                qty_str = str(int(qty_num))
            else:
                qty_str = f"{qty_num:.2f}".rstrip("0").rstrip(".")
        except (ValueError, TypeError):
            qty_str = str(qty)

        unit = str(item.get("unit") or item.get("unidade") or "").strip()
        name = str(
            item.get("name")
            or item.get("description")
            or item.get("descricao")
            or item.get("material_name")
            or item.get("code")
            or ""
        ).strip()

        if not name and not qty:
            continue

        if unit and unit.lower() not in {"un", "unidade", "pc", "peça", "pç"}:
            item_str = f"{qty_str}{unit} {name}".strip()
        else:
            if qty_str and qty_str != "0":
                item_str = f"{qty_str}x {name}".strip()
            else:
                item_str = name

        if item_str:
            formatted_items.append(item_str)

    return ", ".join(formatted_items) if formatted_items else "-"


def format_equipment_readable(equipment_list: Any) -> str:
    """Converte lista de equipamentos em uma string legivel.

    Exemplo: "ONT GPON (SN: 485754431234), DECODER HD (SN: ABC123)"
    """
    if not equipment_list:
        return "-"
    if isinstance(equipment_list, str):
        return equipment_list.strip() or "-"
    if not isinstance(equipment_list, (list, tuple)):
        return "-"

    formatted_items = []
    for item in equipment_list:
        if not isinstance(item, dict):
            continue
        serial = str(
            item.get("serial") or item.get("num_serie") or item.get("identified") or ""
        ).strip()
        name = str(
            item.get("name")
            or item.get("description")
            or item.get("type")
            or item.get("brand")
            or "Equipamento"
        ).strip()

        if serial:
            formatted_items.append(f"{name} (SN: {serial})")
        elif name:
            formatted_items.append(name)

    return ", ".join(formatted_items) if formatted_items else "-"


def build_close_report_xlsx(
    records: list[dict[str, Any]],
    report_date: dt.date,
    profile_label: str = "DOMINIUM",
) -> bytes:
    """Gera o arquivo XLSX formatado do relatorio diario de baixas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Baixas {report_date:%d-%m-%Y}"

    # Habilitar linhas de grade
    ws.views.sheetView[0].showGridLines = True

    # Cores e Estilos
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    meta_font = Font(name="Calibri", size=10, italic=True, color="4B5563")

    status_fills = {
        "confirmed": PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
        "pending": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "uncertain": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
        "failed": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    }
    status_fonts = {
        "confirmed": Font(name="Calibri", size=10, bold=True, color="065F46"),
        "pending": Font(name="Calibri", size=10, bold=True, color="92400E"),
        "uncertain": Font(name="Calibri", size=10, bold=True, color="9A3412"),
        "failed": Font(name="Calibri", size=10, bold=True, color="991B1B"),
    }
    status_labels = {
        "confirmed": "Baixada",
        "pending": "Processando",
        "uncertain": "Incerta",
        "failed": "Falha",
    }

    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )

    # Titulo do Relatorio
    ws.merge_cells("A1:M1")
    ws["A1"] = f"DOMINIUM — RELATÓRIO DIÁRIO DE BAIXAS ({profile_label.upper()})"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(vertical="center")

    now_str = dt.datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ws.merge_cells("A2:M2")
    ws["A2"] = (
        f"Data de Referência: {report_date:%d/%m/%Y} | "
        f"Total de Registros: {len(records)} | Gerado em: {now_str}"
    )
    ws["A2"].font = meta_font
    ws["A2"].alignment = Alignment(vertical="center")

    # Cabecalho da Tabela (Linha 4)
    headers = [
        "Contrato",
        "Técnico",
        "Agenda",
        "Nº OS",
        "Cód. Baixa",
        "Descrição da Baixa",
        "Equipamentos Entrando",
        "Materiais Entrando / Usados",
        "Qtd. Mat.",
        "Horário Baixa",
        "Canal",
        "Resultado",
        "Observação / Retorno",
    ]

    header_row = 4
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num, value=header_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[header_row].height = 28

    # Preenchimento das Linhas
    row_num = 5
    for record in records:
        state = str(record.get("state") or "pending").lower()

        # Formatar Data/Horario
        raw_time = (
            record.get("confirmed_at")
            or record.get("updated_at")
            or record.get("created_at")
            or ""
        )
        formatted_time = "-"
        if raw_time:
            try:
                dt_obj = dt.datetime.fromisoformat(str(raw_time))
                formatted_time = dt_obj.strftime("%H:%M:%S")
            except Exception:
                formatted_time = str(raw_time)[:19]

        # Agenda / Data Agendada
        agenda = str(
            record.get("scheduled_date") or record.get("report_date") or ""
        ).strip()
        if agenda:
            try:
                a_obj = dt.date.fromisoformat(agenda)
                agenda = a_obj.strftime("%d/%m/%Y")
            except Exception:
                pass

        # Materiais e Equipamentos
        mat_summary = (
            record.get("materials_summary")
            or format_materials_readable(record.get("materials"))
        )
        eq_installed = (
            record.get("installed_equipment_summary")
            or format_equipment_readable(record.get("installed_equipment"))
        )

        # Canal
        transport_raw = str(record.get("transport") or "").lower()
        transport_label = "API Imperium" if transport_raw == "official_http" else "DataSnap"

        # Mensagem / Detalhe
        detail = str(
            record.get("message")
            or record.get("detail")
            or record.get("category_label")
            or "-"
        ).strip()

        row_values = [
            str(record.get("contract") or "-"),
            str(record.get("technician") or "-"),
            agenda or "-",
            str(record.get("num_os") or "-"),
            str(record.get("close_code") or "-"),
            str(record.get("close_description") or "-"),
            eq_installed,
            mat_summary,
            int(record.get("material_count") or 0),
            formatted_time,
            transport_label,
            status_labels.get(state, state.upper()),
            detail,
        ]

        for col_num, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.border = thin_border

            # Centralizar codigos, contratos, OS, datas, status
            if col_num in (1, 3, 4, 5, 9, 10, 11, 12):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            # Estilizar coluna de Resultado (Status)
            if col_num == 12:
                if state in status_fills:
                    cell.fill = status_fills[state]
                    cell.font = status_fonts[state]

        # Alternar cor de fundo (zebra striping)
        if row_num % 2 == 0:
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=col_num)
                if col_num != 12:  # Preserva cor do status
                    cell.fill = PatternFill(
                        start_color="F9FAFB", end_color="F9FAFB", fill_type="solid"
                    )

        row_num += 1

    # Ajuste de largura das colunas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:
                continue
            val_str = str(cell.value or "")
            if "\n" in val_str:
                val_str = max(val_str.split("\n"), key=len)
            max_len = max(max_len, len(val_str))

        adjusted_width = min(max(max_len + 4, 12), 45)
        ws.column_dimensions[col_letter].width = adjusted_width

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
