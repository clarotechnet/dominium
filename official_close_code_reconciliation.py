# =============================================================================
# DOMINIUM | MAPA DE RESPONSABILIDADE
#
# IMPERIUM
# - SIM - protocolo, baixa, estoque ou operacao do Imperium.
#
# TOA
# - NAO - este arquivo nao consulta nem automatiza o TOA.
#
# DOMINIUM COMPARTILHADO
# - Apoio local apenas quando necessario ao fluxo Imperium.
#
# Categoria deste arquivo: IMPERIUM.
# Mapa completo: MAPA_DOMINIUM_IMPERIUM_TOA.md
# A ordem executavel abaixo foi preservada para evitar regressao.
# =============================================================================
import argparse
import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from copy import deepcopy
from pathlib import Path

from official_close_code_catalog import (
    build_close_code_catalog,
    extract_pdf_rows,
)


RECONCILIATION_SCHEMA = "dominium_close_code_reconciliation_v1"
REQUIRED_COLUMNS = ("Id", "Código", "Descricao", "Status", "Ativo")


class CloseCodeReconciliationError(ValueError):
    pass


def _sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    try:
        with Path(path).open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
    except OSError as exc:
        raise CloseCodeReconciliationError(
            f"Nao foi possivel ler o arquivo: {exc}"
        ) from exc
    return digest.hexdigest()


def _identifier(value: object, *, field: str) -> str:
    if isinstance(value, bool):
        raise CloseCodeReconciliationError(f"{field} invalido")
    if isinstance(value, int):
        result = str(value)
    elif isinstance(value, float):
        if not value.is_integer():
            raise CloseCodeReconciliationError(f"{field} invalido")
        result = str(int(value))
    else:
        result = str(value or "").strip()
        if re.fullmatch(r"\d+\.0", result):
            result = result[:-2]
    if not re.fullmatch(r"\d+", result):
        raise CloseCodeReconciliationError(f"{field} invalido")
    return result


def normalize_description(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(
        character
        for character in text
        if not unicodedata.combining(character)
    )
    text = re.sub(r"[^A-Za-z0-9]+", " ", text).upper()
    return " ".join(text.split())


def load_registered_close_codes_xlsx(path: str | Path) -> dict:
    source = Path(path).resolve()
    if not source.is_file():
        raise CloseCodeReconciliationError(
            f"Planilha nao encontrada: {source.name}"
        )
    if source.suffix.casefold() != ".xlsx":
        raise CloseCodeReconciliationError(
            "A fonte de codigos cadastrados deve ser XLSX"
        )
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise CloseCodeReconciliationError(
            "openpyxl e necessario para ler o XLSX"
        ) from exc

    workbook = load_workbook(
        source,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        if len(workbook.worksheets) != 1:
            raise CloseCodeReconciliationError(
                "O XLSX deve possuir exatamente uma planilha"
            )
        worksheet = workbook.worksheets[0]
        header_row = None
        header_values = None
        status_legend = ""
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=20),
            start=1,
        ):
            values = tuple(
                str(cell.value or "").strip()
                for cell in row[: len(REQUIRED_COLUMNS)]
            )
            if row_number == 1:
                status_legend = str(row[0].value or "").strip()
            if values == REQUIRED_COLUMNS:
                header_row = row_number
                header_values = values
                break
        if header_row is None or header_values is None:
            raise CloseCodeReconciliationError(
                "Colunas invalidas: esperado Id, Codigo, Descricao, Status e Ativo"
            )

        rows = []
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1),
            start=header_row + 1,
        ):
            values = [cell.value for cell in row[: len(REQUIRED_COLUMNS)]]
            if all(
                value is None or str(value).strip() == ""
                for value in values
            ):
                continue
            if len(values) != len(REQUIRED_COLUMNS):
                raise CloseCodeReconciliationError(
                    f"Linha {row_number} incompleta"
                )
            record_id = _identifier(
                values[0],
                field=f"Id da linha {row_number}",
            )
            code = _identifier(
                values[1],
                field=f"Codigo da linha {row_number}",
            )
            description = str(values[2] or "").strip()
            if not description:
                raise CloseCodeReconciliationError(
                    f"Descricao vazia na linha {row_number}"
                )
            try:
                status = int(values[3])
            except (TypeError, ValueError) as exc:
                raise CloseCodeReconciliationError(
                    f"Status invalido na linha {row_number}"
                ) from exc
            active = str(values[4] or "").strip().upper()
            if active not in {"S", "N"}:
                raise CloseCodeReconciliationError(
                    f"Ativo invalido na linha {row_number}"
                )
            rows.append(
                {
                    "row": row_number,
                    "id": record_id,
                    "code": code,
                    "description": description,
                    "status": status,
                    "active": active,
                }
            )
    finally:
        workbook.close()

    return {
        "source": {
            "file_name": source.name,
            "sha256": _sha256(source),
            "sheet_name": worksheet.title,
            "status_legend": status_legend,
        },
        "rows": rows,
    }


def _pdf_descriptions(entry: dict) -> list[str]:
    if entry.get("status") == "valid":
        return [str(entry.get("description") or "")]
    return [
        str(variant.get("description") or "")
        for variant in entry.get("variants", [])
    ]


def build_close_code_reconciliation(
    pdf_catalog: dict,
    registered_catalog: dict,
) -> dict:
    pdf_snapshot = deepcopy(pdf_catalog)
    registered_snapshot = deepcopy(registered_catalog)
    pdf_entries = pdf_snapshot.get("entries")
    registered_rows = registered_snapshot.get("rows")
    if not isinstance(pdf_entries, list) or not isinstance(
        registered_rows,
        list,
    ):
        raise CloseCodeReconciliationError("Catalogo de entrada invalido")

    pdf_by_code = {}
    for entry in pdf_entries:
        code = _identifier(entry.get("code"), field="Codigo do PDF")
        if code in pdf_by_code:
            raise CloseCodeReconciliationError(
                f"Catalogo PDF possui entrada duplicada para {code}"
            )
        pdf_by_code[code] = entry

    registered_by_code = defaultdict(list)
    for row in registered_rows:
        code = _identifier(row.get("code"), field="Codigo do XLSX")
        registered_by_code[code].append(row)

    pdf_codes = set(pdf_by_code)
    registered_codes = set(registered_by_code)
    missing = [
        deepcopy(pdf_by_code[code])
        for code in sorted(pdf_codes - registered_codes, key=int)
    ]
    registered_only = [
        {
            "code": code,
            "registered_entries": deepcopy(registered_by_code[code]),
        }
        for code in sorted(registered_codes - pdf_codes, key=int)
    ]

    exact_matches = []
    description_mismatches = []
    for code in sorted(pdf_codes & registered_codes, key=int):
        pdf_entry = pdf_by_code[code]
        pdf_descriptions = _pdf_descriptions(pdf_entry)
        registered_entries = registered_by_code[code]
        matches = any(
            normalize_description(row.get("description"))
            == normalize_description(pdf_description)
            for row in registered_entries
            for pdf_description in pdf_descriptions
        )
        result = {
            "code": code,
            "pdf_status": pdf_entry.get("status"),
            "pdf_descriptions": pdf_descriptions,
            "pdf_category": pdf_entry.get("category"),
            "pdf_source_pages": deepcopy(
                pdf_entry.get("source_pages", [])
            ),
            "registered_entries": deepcopy(registered_entries),
        }
        if matches:
            result["status"] = "registered_description_match"
            exact_matches.append(result)
        else:
            result["status"] = "registered_description_differs"
            description_mismatches.append(result)

    registered_duplicates = []
    for code in sorted(registered_by_code, key=int):
        rows = registered_by_code[code]
        if len(rows) < 2:
            continue
        normalized_descriptions = {
            normalize_description(row.get("description"))
            for row in rows
        }
        registered_duplicates.append(
            {
                "code": code,
                "status": (
                    "duplicate_identical"
                    if len(normalized_descriptions) == 1
                    else "duplicate_conflicting"
                ),
                "entries": deepcopy(rows),
            }
        )

    missing_by_category = defaultdict(int)
    for entry in missing:
        category = str(entry.get("category") or "INCONSISTENTE")
        missing_by_category[category] += 1

    summary = {
        "pdf_occurrence_count": int(
            pdf_snapshot.get("occurrence_count", len(pdf_entries))
        ),
        "pdf_unique_code_count": len(pdf_codes),
        "registered_row_count": len(registered_rows),
        "registered_unique_code_count": len(registered_codes),
        "missing_in_imperium_count": len(missing),
        "registered_only_count": len(registered_only),
        "description_match_count": len(exact_matches),
        "description_mismatch_count": len(description_mismatches),
        "registered_duplicate_code_count": len(registered_duplicates),
        "pdf_warning_count": len(pdf_snapshot.get("warnings", [])),
        "missing_by_category": dict(sorted(missing_by_category.items())),
    }
    return {
        "schema": RECONCILIATION_SCHEMA,
        "mode": "offline_read_only",
        "sources": {
            "pdf": deepcopy(pdf_snapshot.get("source", {})),
            "registered_xlsx": deepcopy(
                registered_snapshot.get("source", {})
            ),
        },
        "summary": summary,
        "missing_in_imperium": missing,
        "description_mismatches": description_mismatches,
        "registered_description_matches": exact_matches,
        "registered_only": registered_only,
        "registered_duplicates": registered_duplicates,
        "pdf_warnings": deepcopy(pdf_snapshot.get("warnings", [])),
        "registration_commands_generated": False,
        "registration_executed": False,
        "safety": {
            "network_used": False,
            "imperium_changed": False,
            "payload_generated": False,
        },
    }


def render_reconciliation_text(report: dict) -> str:
    summary = report["summary"]
    lines = [
        "RECONCILIACAO OFFLINE DE CODIGOS DE BAIXA",
        "",
        f"PDF: {report['sources']['pdf'].get('file_name', '')}",
        (
            "XLSX cadastrado: "
            f"{report['sources']['registered_xlsx'].get('file_name', '')}"
        ),
        "",
        "RESUMO",
        f"- Codigos unicos no PDF: {summary['pdf_unique_code_count']}",
        (
            "- Codigos unicos cadastrados: "
            f"{summary['registered_unique_code_count']}"
        ),
        (
            "- Presentes no PDF e ausentes no Imperium: "
            f"{summary['missing_in_imperium_count']}"
        ),
        (
            "- Descricoes divergentes para o mesmo codigo: "
            f"{summary['description_mismatch_count']}"
        ),
        (
            "- Codigos duplicados no XLSX: "
            f"{summary['registered_duplicate_code_count']}"
        ),
        "",
        "FALTANTES NO IMPERIUM",
    ]
    for entry in report["missing_in_imperium"]:
        descriptions = _pdf_descriptions(entry)
        lines.append(
            (
                f"- {entry['code']} | "
                f"{entry.get('category') or 'INCONSISTENTE'} | "
                f"{' / '.join(descriptions)} | "
                f"paginas {entry.get('source_pages', [])}"
            )
        )
    lines.extend(["", "DESCRICOES A VALIDAR"])
    for item in report["description_mismatches"]:
        registered = " / ".join(
            row["description"] for row in item["registered_entries"]
        )
        pdf_descriptions = " / ".join(item["pdf_descriptions"])
        lines.append(
            (
                f"- {item['code']} | XLSX: {registered} | "
                f"PDF: {pdf_descriptions}"
            )
        )
    lines.extend(["", "DUPLICIDADES NO XLSX"])
    for item in report["registered_duplicates"]:
        descriptions = " / ".join(
            row["description"] for row in item["entries"]
        )
        lines.append(
            f"- {item['code']} | {item['status']} | {descriptions}"
        )
    lines.extend(
        [
            "",
            "SEGURANCA",
            "- Nenhum comando de cadastro foi gerado.",
            "- Nenhuma alteracao foi executada no Imperium.",
            "- O relatorio exige validacao humana antes de cadastro.",
            "",
        ]
    )
    return "\n".join(lines)


def _write_json_atomic(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def _write_text_atomic(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(value, encoding="utf-8")
    temporary.replace(path)


def generate_reconciliation(
    *,
    xlsx_path: str | Path,
    pdf_path: str | Path,
    json_output: str | Path,
    text_output: str | Path,
) -> dict:
    pdf_rows, page_count = extract_pdf_rows(pdf_path)
    pdf_catalog = build_close_code_catalog(
        pdf_rows,
        source_file_name=Path(pdf_path).name,
        source_sha256=_sha256(pdf_path),
        page_count=page_count,
    )
    registered_catalog = load_registered_close_codes_xlsx(xlsx_path)
    report = build_close_code_reconciliation(
        pdf_catalog,
        registered_catalog,
    )
    _write_json_atomic(Path(json_output), report)
    _write_text_atomic(
        Path(text_output),
        render_reconciliation_text(report),
    )
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Compara offline os codigos esperados no PDF com os "
            "codigos cadastrados no XLSX."
        ),
    )
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--pdf", required=True)
    parser.add_argument("--json-output", required=True)
    parser.add_argument("--text-output", required=True)
    args = parser.parse_args(argv)
    report = generate_reconciliation(
        xlsx_path=args.xlsx,
        pdf_path=args.pdf,
        json_output=args.json_output,
        text_output=args.text_output,
    )
    print(json.dumps(report["summary"], ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
