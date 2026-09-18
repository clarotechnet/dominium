import hashlib
import importlib
import json
import socket
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from official_material_catalog import (
    CatalogSchemaError,
    CatalogValidationError,
    build_contract_2221170_report,
    load_xlsx_catalog,
)


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class OfficialMaterialCatalogTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)

    def create_workbook(
        self,
        rows: list[tuple[object, object, object]],
        *,
        headers: tuple[str, ...] = ("Id", "Código", "Descrição"),
        name: str = "catalog.xlsx",
    ) -> Path:
        path = self.directory / name
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def load(
        self,
        rows: list[tuple[object, object, object]],
    ):
        return load_xlsx_catalog(
            self.create_workbook(rows),
            imported_at="2026-07-23T12:00:00-03:00",
        )

    def test_reads_valid_workbook_and_builds_indexes(self) -> None:
        catalog = self.load(
            [
                (1, "00123456", "MATERIAL ALFA"),
                (2, 22069613, "CONECTOR FO CAMPO FAST SC APC"),
            ]
        )

        self.assertEqual(catalog.metadata["record_count"], 2)
        self.assertEqual(catalog.metadata["unique_code_count"], 2)
        self.assertEqual(catalog.lookup_code("00123456")["ids"], ["1"])
        self.assertEqual(
            catalog.lookup_description("MATERIAL ALFA")["selected_code"],
            "00123456",
        )

    def test_numeric_code_respects_excel_leading_zero_format(self) -> None:
        path = self.create_workbook(
            [(1, 1234567, "MATERIAL COM ZERO")]
        )
        workbook = load_workbook(path)
        workbook.active["B2"].number_format = "00000000"
        workbook.save(path)
        workbook.close()

        catalog = load_xlsx_catalog(path)

        self.assertIsNotNone(catalog.lookup_code("01234567"))

    def test_missing_columns_are_rejected(self) -> None:
        path = self.create_workbook(
            [(1, "22069613", "CONECTOR")],
            headers=("Id", "Código", "Nome"),
        )

        with self.assertRaisesRegex(CatalogSchemaError, "Colunas invalidas"):
            load_xlsx_catalog(path)

    def test_duplicate_code_with_identical_description_is_registered(self) -> None:
        catalog = self.load(
            [
                (1, "22069613", "CONECTOR FO CAMPO FAST SC APC"),
                (2, "22069613", "CONECTOR FO CAMPO FAST SC/APC"),
            ]
        )

        self.assertEqual(len(catalog.duplicate_codes), 1)
        self.assertEqual(
            catalog.duplicate_codes[0]["kind"],
            "identical_description",
        )
        self.assertEqual(catalog.critical_errors, [])

    def test_duplicate_code_with_divergent_descriptions_is_critical(self) -> None:
        catalog = self.load(
            [
                (1, "22055857", "MARCADOR CASA PTO NR 8"),
                (2, "22055857", "OUTRO MATERIAL"),
            ]
        )

        self.assertEqual(len(catalog.critical_errors), 1)
        self.assertEqual(
            catalog.critical_errors[0]["type"],
            "duplicate_code_with_divergent_descriptions",
        )

    def test_duplicate_description_returns_ambiguity_without_selection(self) -> None:
        catalog = self.load(
            [
                (1, "22057620", "CONECTOR FO CAMPO FAST SC APC"),
                (2, "22069613", "CONECTOR FO CAMPO FAST SC/APC"),
            ]
        )

        result = catalog.lookup_description(
            "CONECTOR FO CAMPO FAST SC-APC"
        )

        self.assertTrue(result["ambiguous"])
        self.assertEqual(result["status"], "ambiguous")
        self.assertIsNone(result["selected_code"])
        self.assertEqual(result["codes"], ["22057620", "22069613"])

    def test_exact_code_and_description_are_validated(self) -> None:
        catalog = self.load(
            [(1, "22069613", "CONECTOR FO CAMPO FAST SC APC")]
        )

        result = catalog.resolve_evidence(
            toa_code="22069613",
            toa_description="22069613_CONECTOR FO CAMPO FAST SC/APC",
            desktop_code="22069613",
            desktop_description="CONECTOR FO CAMPO FAST SC APC",
        )

        self.assertEqual(result["status"], "validado")
        self.assertEqual(result["official"]["description"], "CONECTOR FO CAMPO FAST SC APC")

    def test_punctuation_is_normalized_for_exact_comparison(self) -> None:
        catalog = self.load(
            [(1, "22069613", "CONECTOR FO CAMPO FAST SC APC")]
        )

        lookup = catalog.lookup_description(
            "22069613_CONECTOR FO CAMPO FAST SC/APC",
            "22069613",
        )

        self.assertEqual(lookup["status"], "matched")
        self.assertEqual(lookup["selected_code"], "22069613")

    def test_fuzzy_description_is_not_used(self) -> None:
        catalog = self.load(
            [(1, "22069613", "CONECTOR FO CAMPO FAST SC APC")]
        )

        result = catalog.lookup_description("CONECTOR CAMPO FAST")

        self.assertEqual(result["status"], "not_found")
        self.assertEqual(result["candidates"], [])

    def test_incompatible_code_and_description_are_blocked(self) -> None:
        catalog = self.load(
            [
                (1, "22056408", "CONECTOR ATENUADOR 06DB"),
                (2, "22064608", "FITA ISOLANTE 3M HIGHLAND 19MM X 20M"),
            ]
        )

        result = catalog.resolve_evidence(
            toa_code="22056408",
            toa_description="FITA ISOLANTE 3M HIGHLAND 19MM X 20M",
            desktop_code="22056408",
            desktop_description="FITA ISOLANTE 3M HIGHLAND 19MM X 20M",
        )

        self.assertEqual(result["status"], "bloqueado")
        self.assertIn(
            "toa_code_and_description_incompatible",
            result["reasons"],
        )

    def test_unrelated_divergent_catalog_entry_does_not_block_payload(self) -> None:
        catalog = self.load(
            [
                (1, "22055857", "MARCADOR CASA PTO NR 8"),
                (2, "22055857", "OUTRO MATERIAL"),
                (3, "22069613", "CONECTOR FO CAMPO FAST SC APC"),
            ]
        )
        review = catalog.resolve_evidence(
            toa_code="22069613",
            toa_description="CONECTOR FO CAMPO FAST SC APC",
            desktop_code="22069613",
            desktop_description="CONECTOR FO CAMPO FAST SC APC",
        )

        catalog.assert_payload_catalog_safe(
            [review],
            official_stock_proven=True,
        )
        self.assertEqual(catalog.critical_errors_for_reviews([review]), [])

    def test_used_divergent_catalog_code_blocks_payload(self) -> None:
        catalog = self.load(
            [
                (1, "22055857", "MARCADOR CASA PTO NR 8"),
                (2, "22055857", "OUTRO MATERIAL"),
            ]
        )
        review = catalog.resolve_evidence(
            toa_code="22055857",
            toa_description="MARCADOR CASA PTO NR 8",
            desktop_code="22055857",
            desktop_description="MARCADOR CASA PTO NR 8",
        )

        with self.assertRaisesRegex(
            CatalogValidationError,
            "dependem de codigo",
        ):
            catalog.assert_payload_catalog_safe(
                [review],
                official_stock_proven=True,
            )

    def test_candidate_from_divergent_entry_blocks_payload(self) -> None:
        catalog = self.load(
            [
                (1, "22055857", "MARCADOR CASA PTO NR 8"),
                (2, "22055857", "OUTRO MATERIAL"),
                (3, "22099999", "MARCADOR CASA PTO NR 8"),
            ]
        )
        review = catalog.resolve_evidence(
            toa_code="22099999",
            toa_description="MARCADOR CASA PTO NR 8",
            desktop_code="22099999",
            desktop_description="MARCADOR CASA PTO NR 8",
        )

        with self.assertRaisesRegex(
            CatalogValidationError,
            "dependem de codigo",
        ):
            catalog.assert_payload_catalog_safe(
                [review],
                official_stock_proven=True,
            )

    def test_description_candidate_never_creates_automatic_alias(self) -> None:
        catalog = self.load(
            [
                (1, "22025072", "FITA ISOLANTE 3M 33+"),
                (2, "22056408", "CONECTOR ATENUADOR 06DB"),
                (3, "22064608", "FITA ISOLANTE 3M HIGHLAND 19MM X 20M"),
            ]
        )

        result = catalog.resolve_evidence(
            toa_code="22025072",
            toa_description="22025072_FITA ISOLANTE 3M 33+",
            desktop_code="22064608",
            desktop_description="FITA ISOLANTE 3M HIGHLAND 19MM X 20M",
        )

        self.assertEqual(result["status"], "bloqueado")
        self.assertIn(
            "desktop_code_differs_no_automatic_alias",
            result["reasons"],
        )
        self.assertNotIn(
            "22056408",
            [
                candidate["code"]
                for candidate in result[
                    "desktop_candidates_by_exact_normalized_description"
                ]
            ],
        )

    def test_contract_report_keeps_stock_and_group_as_blockers(self) -> None:
        rows = [
            (index, code, description)
            for index, (code, description) in enumerate(
                [
                    ("22061736", "CABO DROP 1FO LOW F FIG8 LOW CINZA"),
                    ("22069613", "CONECTOR FO CAMPO FAST SC APC"),
                    ("22025072", "FITA ISOLANTE 3M 33+"),
                    ("22056332", "FITA AUTO-FUSAO 23LB 19X10MM 3M NET"),
                    ("22056343", "MARCADOR CASA PTO NR 1"),
                    ("22056341", "MARCADOR CASA PTO NR 7"),
                    ("22056344", "MARCADOR CASA PTO NR 2"),
                    ("22057659", "ESTICADOR CUNHA P/DROP FO SDA1 DPR"),
                    ("22056408", "CONECTOR ATENUADOR 06DB"),
                    ("22064608", "FITA ISOLANTE 3M HIGHLAND 19MM X 20M"),
                ],
                start=1,
            )
        ]
        report = build_contract_2221170_report(self.load(rows))

        self.assertFalse(report["payload_catalog_validation"]["authorized"])
        self.assertIn(
            "official_stock_not_proven:Z637677:installer_id_328898",
            report["blockers"],
        )
        self.assertIn(
            "catalog_has_no_stock_or_group_data",
            report["blockers"],
        )
        self.assertEqual(
            report["highlighted_findings"]["22056408"]["descriptions"],
            ["CONECTOR ATENUADOR 06DB"],
        )
        self.assertEqual(
            report["highlighted_findings"]["22064608"]["classification"],
            "description_candidate_only",
        )

    def test_contract_report_keeps_unrelated_catalog_error_as_warning(self) -> None:
        rows = [
            (index, code, description)
            for index, (code, description) in enumerate(
                [
                    ("22061736", "CABO DROP 1FO LOW F FIG8 LOW CINZA"),
                    ("22069613", "CONECTOR FO CAMPO FAST SC APC"),
                    ("22025072", "FITA ISOLANTE 3M 33+"),
                    ("22056332", "FITA AUTO-FUSAO 23LB 19X10MM 3M NET"),
                    ("22056343", "MARCADOR CASA PTO NR 1"),
                    ("22056341", "MARCADOR CASA PTO NR 7"),
                    ("22056344", "MARCADOR CASA PTO NR 2"),
                    ("22057659", "ESTICADOR CUNHA P/DROP FO SDA1 DPR"),
                    ("22064608", "FITA ISOLANTE 3M HIGHLAND 19MM X 20M"),
                    ("22055857", "MATERIAL CONFLITANTE UM"),
                    ("22055857", "MATERIAL CONFLITANTE DOIS"),
                ],
                start=1,
            )
        ]

        report = build_contract_2221170_report(self.load(rows))

        self.assertNotIn(
            "related_official_catalog_critical_errors:1",
            report["blockers"],
        )
        self.assertIn(
            "unrelated_official_catalog_critical_errors:1",
            report["warnings"],
        )
        self.assertEqual(len(report["catalog"]["critical_errors"]), 1)
        self.assertEqual(report["catalog"]["related_critical_errors"], [])

    def test_import_is_read_only_and_uses_no_network(self) -> None:
        path = self.create_workbook(
            [(1, "22069613", "CONECTOR FO CAMPO FAST SC APC")]
        )
        before_bytes = path.read_bytes()
        before_hash = file_hash(path)
        before_stat = path.stat()

        with patch.object(
            socket,
            "create_connection",
            side_effect=AssertionError("network forbidden"),
        ), patch.object(
            socket,
            "socket",
            side_effect=AssertionError("network forbidden"),
        ):
            catalog = load_xlsx_catalog(path)

        self.assertEqual(catalog.metadata["record_count"], 1)
        self.assertEqual(path.read_bytes(), before_bytes)
        self.assertEqual(file_hash(path), before_hash)
        self.assertEqual(path.stat().st_mtime_ns, before_stat.st_mtime_ns)

    def test_importing_module_has_no_network_side_effect(self) -> None:
        module_name = "official_material_catalog"
        original = sys.modules.pop(module_name, None)
        self.addCleanup(
            lambda: sys.modules.__setitem__(module_name, original)
            if original is not None
            else None
        )

        with patch.object(
            socket,
            "create_connection",
            side_effect=AssertionError("network forbidden"),
        ), patch.object(
            socket,
            "socket",
            side_effect=AssertionError("network forbidden"),
        ):
            imported = importlib.import_module(module_name)

        self.assertTrue(hasattr(imported, "load_xlsx_catalog"))

    def test_generated_index_contains_no_aliases_or_source_path(self) -> None:
        catalog = self.load(
            [(1, "22069613", "CONECTOR FO CAMPO FAST SC APC")]
        )

        serialized = json.dumps(catalog.to_index_dict(), ensure_ascii=False)

        self.assertEqual(catalog.to_index_dict()["aliases"], [])
        self.assertNotIn(str(self.directory), serialized)


if __name__ == "__main__":
    unittest.main()
