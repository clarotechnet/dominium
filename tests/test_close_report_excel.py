from __future__ import annotations

import datetime as dt
import io
import unittest
import openpyxl

from close_report_excel import (
    build_close_report_xlsx,
    format_equipment_readable,
    format_materials_readable,
)


class CloseReportExcelTests(unittest.TestCase):
    def test_format_materials_readable_with_list(self) -> None:
        materials = [
            {"name": "CABO 2FO", "quantity": 100, "unit": "m"},
            {"name": "FITA ISOLANTE", "quantity": 6, "unit": "m"},
            {"name": "CONECTOR REUTILIZAVEL", "quantity": 2, "unit": "UN"},
        ]
        result = format_materials_readable(materials)
        self.assertEqual(result, "100m CABO 2FO, 6m FITA ISOLANTE, 2x CONECTOR REUTILIZAVEL")

    def test_format_materials_readable_empty(self) -> None:
        self.assertEqual(format_materials_readable([]), "-")
        self.assertEqual(format_materials_readable(None), "-")

    def test_format_equipment_readable_with_list(self) -> None:
        equipment = [
            {"name": "ONT GPON", "serial": "4857544312345678"},
            {"name": "DECODER HD", "serial": "ABC123456"},
        ]
        result = format_equipment_readable(equipment)
        self.assertEqual(result, "ONT GPON (SN: 4857544312345678), DECODER HD (SN: ABC123456)")

    def test_build_close_report_xlsx_structure(self) -> None:
        records = [
            {
                "contract": "123456",
                "technician": "FERNANDO ANTONIO",
                "scheduled_date": "2026-08-25",
                "num_os": "2652991103",
                "close_code": "409",
                "close_description": "INSTALACAO CONCLUIDA",
                "installed_equipment": [{"name": "ONT GPON", "serial": "485754431234"}],
                "materials": [
                    {"name": "CABO 2FO", "quantity": 100, "unit": "m"},
                    {"name": "FITA ISOLANTE", "quantity": 6, "unit": "m"},
                ],
                "material_count": 2,
                "confirmed_at": "2026-08-25T14:32:05-03:00",
                "transport": "official_http",
                "state": "confirmed",
                "message": "Baixa confirmada no Imperium",
            }
        ]
        today = dt.date(2026, 8, 25)
        raw_bytes = build_close_report_xlsx(records, today, "NATAL")
        self.assertGreater(len(raw_bytes), 1000)

        # Validar leitura do workbook
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))
        ws = wb.active
        self.assertEqual(ws.title, "Baixas 25-08-2026")
        self.assertIn("DOMINIUM", str(ws["A1"].value))

        # Linha 4: Cabeçalho
        self.assertEqual(ws.cell(row=4, column=1).value, "Contrato")
        self.assertEqual(ws.cell(row=4, column=8).value, "Materiais Entrando / Usados")

        # Linha 5: Dados
        self.assertEqual(ws.cell(row=5, column=1).value, "123456")
        self.assertEqual(ws.cell(row=5, column=2).value, "FERNANDO ANTONIO")
        self.assertEqual(ws.cell(row=5, column=8).value, "100m CABO 2FO, 6m FITA ISOLANTE")
        self.assertEqual(ws.cell(row=5, column=12).value, "Baixada")


if __name__ == "__main__":
    unittest.main()
