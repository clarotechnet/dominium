# =============================================================================
# DOMINIUM | MAPA DE RESPONSABILIDADE
#
# IMPERIUM
# - SIM - protocolo, baixa, estoque ou operacao do Imperium.
#
# TOA
# - NAO - este arquivo nao consulta nem automatiza o TOA.
#
# DOMINIUM COMPARTILHADO
# - Apoio local apenas quando necessario ao fluxo Imperium.
#
# Categoria deste arquivo: IMPERIUM.
# Mapa completo: MAPA_DOMINIUM_IMPERIUM_TOA.md
# A ordem executavel abaixo foi preservada para evitar regressao.
# =============================================================================
import argparse
import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable


REQUIRED_COLUMNS = ("Id", "Código", "Descrição")
CATALOG_SCHEMA = "dominium_official_material_catalog_v1"
REPORT_SCHEMA = "dominium_official_material_catalog_review_v1"


class CatalogError(ValueError):
    pass


class CatalogSchemaError(CatalogError):
    pass


class CatalogValidationError(CatalogError):
    pass


@dataclass(frozen=True)
class CatalogRecord:
    record_id: str
    code: str
    description: str
    row_number: int

    def to_dict(self) -> dict:
        return {
            "id": self.record_id,
            "code": self.code,
            "description": self.description,
            "row_number": self.row_number,
        }


def normalize_description(value: object, code: object = "") -> str:
    text = " ".join(str(value or "").strip().split())
    normalized_code = normalize_identifier(code, "Codigo", allow_empty=True)
    if normalized_code:
        text = re.sub(
            rf"^\s*{re.escape(normalized_code)}(?=$|[\s_/\-])[\s_/\-]*",
            "",
            text,
            flags=re.IGNORECASE,
        )
    ascii_text = (
        unicodedata.normalize("NFD", text)
        .encode("ascii", "ignore")
        .decode("ascii")
        .upper()
    )
    return " ".join(re.findall(r"[A-Z0-9]+", ascii_text))


def normalize_identifier(
    value: object,
    label: str,
    *,
    allow_empty: bool = False,
    number_format: str = "General",
) -> str:
    if value is None:
        if allow_empty:
            return ""
        raise CatalogSchemaError(f"{label} vazio")
    if isinstance(value, bool):
        raise CatalogSchemaError(f"{label} invalido")
    if isinstance(value, str):
        result = value.strip()
    else:
        try:
            numeric = Decimal(str(value))
        except (InvalidOperation, ValueError) as exc:
            raise CatalogSchemaError(f"{label} invalido") from exc
        if not numeric.is_finite() or numeric != numeric.to_integral_value():
            raise CatalogSchemaError(f"{label} deve ser inteiro ou texto")
        result = format(numeric.quantize(Decimal("1")), "f")
        zero_format = re.fullmatch(r"0+", str(number_format or ""))
        if zero_format:
            result = result.zfill(len(zero_format.group(0)))
    if not result and not allow_empty:
        raise CatalogSchemaError(f"{label} vazio")
    return result


def _cell_identifier(cell: object, label: str) -> str:
    return normalize_identifier(
        getattr(cell, "value", None),
        label,
        number_format=str(getattr(cell, "number_format", "General")),
    )


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


class OfficialMaterialCatalog:
    def __init__(
        self,
        *,
        source_name: str,
        source_sha256: str,
        imported_at: str,
        records: Iterable[CatalogRecord],
    ) -> None:
        self.source_name = str(source_name)
        self.source_sha256 = str(source_sha256).lower()
        self.imported_at = str(imported_at)
        self.records = tuple(records)

        records_by_code: dict[str, list[CatalogRecord]] = defaultdict(list)
        codes_by_description: dict[str, set[str]] = defaultdict(set)
        for record in self.records:
            records_by_code[record.code].append(record)
            codes_by_description[
                normalize_description(record.description)
            ].add(record.code)

        self.duplicate_codes = []
        self.critical_errors = []
        self.by_code = {}
        for code in sorted(records_by_code):
            code_records = records_by_code[code]
            descriptions_by_normalized: dict[str, set[str]] = defaultdict(set)
            for record in code_records:
                descriptions_by_normalized[
                    normalize_description(record.description)
                ].add(record.description)
            descriptions = sorted(
                {
                    description
                    for values in descriptions_by_normalized.values()
                    for description in values
                }
            )
            entry = {
                "code": code,
                "ids": sorted({record.record_id for record in code_records}),
                "descriptions": descriptions,
                "rows": sorted(record.row_number for record in code_records),
            }
            self.by_code[code] = entry
            if len(code_records) > 1:
                duplicate = {
                    **entry,
                    "kind": (
                        "identical_description"
                        if len(descriptions_by_normalized) == 1
                        else "divergent_description"
                    ),
                }
                self.duplicate_codes.append(duplicate)
                if len(descriptions_by_normalized) > 1:
                    self.critical_errors.append(
                        {
                            "type": "duplicate_code_with_divergent_descriptions",
                            **entry,
                        }
                    )

        self.by_description = {
            description: sorted(codes)
            for description, codes in sorted(codes_by_description.items())
        }
        self.ambiguous_descriptions = [
            {
                "normalized_description": description,
                "codes": codes,
            }
            for description, codes in self.by_description.items()
            if len(codes) > 1
        ]

    @property
    def metadata(self) -> dict:
        return {
            "schema": CATALOG_SCHEMA,
            "source_file": self.source_name,
            "source_sha256": self.source_sha256,
            "imported_at": self.imported_at,
            "required_columns": list(REQUIRED_COLUMNS),
            "record_count": len(self.records),
            "unique_code_count": len(self.by_code),
            "duplicate_code_count": len(self.duplicate_codes),
            "critical_error_count": len(self.critical_errors),
            "ambiguous_description_count": len(self.ambiguous_descriptions),
        }

    def lookup_code(self, code: object) -> dict | None:
        normalized_code = normalize_identifier(
            code,
            "Codigo",
            allow_empty=True,
        )
        entry = self.by_code.get(normalized_code)
        return json.loads(json.dumps(entry)) if entry is not None else None

    def lookup_description(self, description: object, code: object = "") -> dict:
        normalized = normalize_description(description, code)
        codes = list(self.by_description.get(normalized, []))
        return {
            "normalized_description": normalized,
            "codes": codes,
            "candidates": [
                json.loads(json.dumps(self.by_code[candidate]))
                for candidate in codes
            ],
            "ambiguous": len(codes) > 1,
            "status": (
                "not_found"
                if not codes
                else "ambiguous"
                if len(codes) > 1
                else "matched"
            ),
            "selected_code": codes[0] if len(codes) == 1 else None,
        }

    def resolve_evidence(
        self,
        *,
        toa_code: object,
        toa_description: object,
        desktop_code: object,
        desktop_description: object,
    ) -> dict:
        code = normalize_identifier(toa_code, "Codigo TOA")
        description = str(toa_description or "").strip()
        desktop_normalized_code = normalize_identifier(
            desktop_code,
            "Codigo desktop",
        )
        desktop_description_text = str(desktop_description or "").strip()

        code_entry = self.lookup_code(code)
        description_lookup = self.lookup_description(description, code)
        desktop_entry = self.lookup_code(desktop_normalized_code)
        desktop_lookup = self.lookup_description(
            desktop_description_text,
            desktop_normalized_code,
        )

        evidence = []
        status = "validado"
        reasons = []

        if code_entry is None:
            status = (
                "ambiguo"
                if description_lookup["ambiguous"]
                else "bloqueado"
            )
            reasons.append("toa_code_not_found_in_official_catalog")
        else:
            official_descriptions = {
                normalize_description(value)
                for value in code_entry["descriptions"]
            }
            if normalize_description(description, code) not in official_descriptions:
                status = "bloqueado"
                reasons.append("toa_code_and_description_incompatible")
            else:
                evidence.append("toa_code_and_description_match_official_catalog")

        if desktop_entry is None:
            status = "bloqueado"
            reasons.append("desktop_code_not_found_in_official_catalog")
        else:
            desktop_descriptions = {
                normalize_description(value)
                for value in desktop_entry["descriptions"]
            }
            if (
                normalize_description(
                    desktop_description_text,
                    desktop_normalized_code,
                )
                not in desktop_descriptions
            ):
                status = "bloqueado"
                reasons.append("desktop_code_and_description_incompatible")
            else:
                evidence.append(
                    "desktop_code_and_description_match_official_catalog"
                )

        if desktop_normalized_code != code:
            if desktop_lookup["ambiguous"]:
                status = "ambiguo"
                reasons.append("desktop_description_has_multiple_official_codes")
            else:
                status = "bloqueado"
                reasons.append("desktop_code_differs_no_automatic_alias")
            evidence.append("description_candidate_is_not_send_authorization")

        if status == "validado" and description_lookup["ambiguous"]:
            evidence.append(
                "description_only_lookup_is_ambiguous_but_exact_code_is_primary"
            )

        return {
            "toa": {
                "code": code,
                "description": description,
            },
            "desktop": {
                "code": desktop_normalized_code,
                "description": desktop_description_text,
            },
            "official": {
                "code_entry": code_entry,
                "description": (
                    code_entry["descriptions"][0]
                    if code_entry and len(code_entry["descriptions"]) == 1
                    else None
                ),
            },
            "official_candidates_by_exact_normalized_description": (
                description_lookup["candidates"]
            ),
            "desktop_candidates_by_exact_normalized_description": (
                desktop_lookup["candidates"]
            ),
            "ambiguity": {
                "toa_description": description_lookup["ambiguous"],
                "desktop_description": desktop_lookup["ambiguous"],
                "blocks_description_only_resolution": (
                    description_lookup["ambiguous"]
                    or desktop_lookup["ambiguous"]
                ),
            },
            "evidence": evidence,
            "status": status,
            "reasons": reasons,
        }

    def assert_payload_catalog_safe(
        self,
        material_reviews: Iterable[dict],
        *,
        official_stock_proven: bool,
    ) -> None:
        reviews = list(material_reviews)
        related_critical_errors = self.critical_errors_for_reviews(reviews)
        if related_critical_errors:
            raise CatalogValidationError(
                "Materiais dependem de codigo do catalogo com descricoes "
                "divergentes"
            )
        blocked = [
            review
            for review in reviews
            if review.get("status") != "validado"
        ]
        if blocked:
            raise CatalogValidationError(
                "Catalogo nao validou todos os materiais do payload"
            )
        if not official_stock_proven:
            raise CatalogValidationError(
                "Estoque oficial do tecnico nao foi comprovado"
            )

    def critical_errors_for_reviews(
        self,
        material_reviews: Iterable[dict],
    ) -> list[dict]:
        dependency_codes: set[str] = set()
        dependency_descriptions: set[str] = set()

        def add_code(value: object) -> None:
            code = normalize_identifier(
                value,
                "Codigo",
                allow_empty=True,
            )
            if code:
                dependency_codes.add(code)

        def add_description(value: object, code: object = "") -> None:
            description = normalize_description(value, code)
            if description:
                dependency_descriptions.add(description)

        for review in material_reviews:
            if not isinstance(review, dict):
                continue
            for section_name in ("toa", "desktop"):
                section = review.get(section_name, {})
                if not isinstance(section, dict):
                    continue
                add_code(section.get("code", ""))
                add_description(
                    section.get("description", ""),
                    section.get("code", ""),
                )

            official = review.get("official", {})
            if isinstance(official, dict):
                code_entry = official.get("code_entry")
                if isinstance(code_entry, dict):
                    add_code(code_entry.get("code", ""))
                    for description in code_entry.get("descriptions", []):
                        add_description(description)

            for candidate_key in (
                "official_candidates_by_exact_normalized_description",
                "desktop_candidates_by_exact_normalized_description",
            ):
                candidates = review.get(candidate_key, [])
                if not isinstance(candidates, list):
                    continue
                for candidate in candidates:
                    if not isinstance(candidate, dict):
                        continue
                    add_code(candidate.get("code", ""))
                    for description in candidate.get("descriptions", []):
                        add_description(description)

        related = []
        for error in self.critical_errors:
            error_code = str(error.get("code", "")).strip()
            error_descriptions = {
                normalize_description(description)
                for description in error.get("descriptions", [])
                if normalize_description(description)
            }
            if (
                error_code in dependency_codes
                or bool(error_descriptions & dependency_descriptions)
            ):
                related.append(json.loads(json.dumps(error)))
        return related

    def to_index_dict(self) -> dict:
        return {
            "metadata": self.metadata,
            "by_code": self.by_code,
            "by_exact_normalized_description": self.by_description,
            "duplicate_codes": self.duplicate_codes,
            "critical_errors": self.critical_errors,
            "ambiguities": self.ambiguous_descriptions,
            "aliases": [],
            "contains_stock": False,
            "contains_groups": False,
        }


def load_xlsx_catalog(
    path: str | Path,
    *,
    imported_at: str | None = None,
) -> OfficialMaterialCatalog:
    source_path = Path(path).resolve()
    if not source_path.is_file():
        raise CatalogSchemaError(f"Planilha nao encontrada: {source_path.name}")
    if source_path.suffix.casefold() != ".xlsx":
        raise CatalogSchemaError("Catalogo deve ser um arquivo XLSX")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise CatalogSchemaError(
            "openpyxl e necessario para ler o catalogo XLSX"
        ) from exc

    workbook = load_workbook(
        source_path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        if len(workbook.worksheets) != 1:
            raise CatalogSchemaError(
                "Catalogo deve possuir exatamente uma planilha"
            )
        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows()
        try:
            header_cells = next(rows)
        except StopIteration as exc:
            raise CatalogSchemaError("Catalogo vazio") from exc
        headers = tuple(
            str(getattr(cell, "value", "") or "").strip()
            for cell in header_cells
        )
        if headers != REQUIRED_COLUMNS:
            raise CatalogSchemaError(
                "Colunas invalidas: esperado Id, Codigo e Descricao"
            )

        records = []
        for row_number, row in enumerate(rows, start=2):
            values = [getattr(cell, "value", None) for cell in row]
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            if len(row) != 3:
                raise CatalogSchemaError(
                    f"Linha {row_number} possui quantidade invalida de colunas"
                )
            record_id = _cell_identifier(row[0], f"Id da linha {row_number}")
            code = _cell_identifier(row[1], f"Codigo da linha {row_number}")
            description = str(row[2].value or "").strip()
            if not description:
                raise CatalogSchemaError(
                    f"Descricao vazia na linha {row_number}"
                )
            records.append(
                CatalogRecord(
                    record_id=record_id,
                    code=code,
                    description=description,
                    row_number=row_number,
                )
            )
    finally:
        workbook.close()

    return OfficialMaterialCatalog(
        source_name=source_path.name,
        source_sha256=_sha256_file(source_path),
        imported_at=imported_at or _now_iso(),
        records=records,
    )


CONTRACT_2221170_MATERIAL_EVIDENCE = (
    {
        "toa_code": "22061736",
        "toa_description": "22061736_CABO DROP 1FO LOW F FIG8 LOW CINZA",
        "desktop_code": "22061736",
        "desktop_description": "CABO DROP 1FO LOW F FIG8 LOW CINZA",
    },
    {
        "toa_code": "22069613",
        "toa_description": "22069613_CONECTOR FO CAMPO FAST SC/APC",
        "desktop_code": "22069613",
        "desktop_description": "CONECTOR FO CAMPO FAST SC APC",
    },
    {
        "toa_code": "22025072",
        "toa_description": "22025072_FITA ISOLANTE 3M 33+",
        "desktop_code": "22064608",
        "desktop_description": "FITA ISOLANTE 3M HIGHLAND 19MM X 20M",
    },
    {
        "toa_code": "22056332",
        "toa_description": "22056332_FITA AUTO-FUSAO 23LB 19X10MM 3M NET",
        "desktop_code": "22056332",
        "desktop_description": "FITA AUTO-FUSAO 23LB 19X10MM 3M NET",
    },
    {
        "toa_code": "22056343",
        "toa_description": "22056343_MARCADOR CASA PTO NR 1",
        "desktop_code": "22056343",
        "desktop_description": "MARCADOR CASA PTO NR 1",
    },
    {
        "toa_code": "22056341",
        "toa_description": "22056341_MARCADOR CASA PTO NR 7",
        "desktop_code": "22056341",
        "desktop_description": "MARCADOR CASA PTO NR 7",
    },
    {
        "toa_code": "22056344",
        "toa_description": "22056344_MARCADOR CASA PTO NR 2",
        "desktop_code": "22056344",
        "desktop_description": "MARCADOR CASA PTO NR 2",
    },
    {
        "toa_code": "22057659",
        "toa_description": "22057659_ESTICADOR CUNHA P/DROP FO SDA1 DPR",
        "desktop_code": "22057659",
        "desktop_description": "ESTICADOR CUNHA P/DROP FO SDA1 DPR",
    },
)


def build_contract_2221170_report(
    catalog: OfficialMaterialCatalog,
) -> dict:
    reviews = [
        catalog.resolve_evidence(**item)
        for item in CONTRACT_2221170_MATERIAL_EVIDENCE
    ]
    blockers = []
    related_critical_errors = catalog.critical_errors_for_reviews(reviews)
    if related_critical_errors:
        blockers.append(
            "related_official_catalog_critical_errors:"
            f"{len(related_critical_errors)}"
        )
    blockers.extend(
        f"material_catalog_{review['status']}:{review['toa']['code']}"
        for review in reviews
        if review["status"] != "validado"
    )
    blockers.extend(
        [
            "official_stock_not_proven:Z637677:installer_id_328898",
            "catalog_has_no_stock_or_group_data",
        ]
    )

    return {
        "schema": REPORT_SCHEMA,
        "generated_at": _now_iso(),
        "mode": "offline_catalog_review",
        "contract": "2221170",
        "profile": "TECHNET NATAL",
        "catalog": {
            "metadata": catalog.metadata,
            "critical_errors": catalog.critical_errors,
            "related_critical_errors": related_critical_errors,
            "duplicate_codes": catalog.duplicate_codes,
        },
        "materials": reviews,
        "highlighted_findings": {
            "22056408": catalog.lookup_code("22056408"),
            "22064608": {
                "catalog_entry": catalog.lookup_code("22064608"),
                "exact_description_lookup": catalog.lookup_description(
                    "FITA ISOLANTE 3M HIGHLAND 19MM X 20M"
                ),
                "classification": "description_candidate_only",
                "stock_proven_for_technician": False,
                "authorized_for_send": False,
            },
            "invalid_mapping_removed": {
                "from": "22025072",
                "to": "22056408",
                "reason": (
                    "22056408 is CONECTOR ATENUADOR 06DB in the "
                    "authoritative official catalog"
                ),
            },
        },
        "limitations": {
            "contains_stock": False,
            "contains_groups": False,
            "can_authorize_distribution_alone": False,
            "can_authorize_send_alone": False,
            "aliases_created": [],
        },
        "blockers": blockers,
        "warnings": (
            [
                "unrelated_official_catalog_critical_errors:"
                f"{len(catalog.critical_errors) - len(related_critical_errors)}"
            ]
            if len(catalog.critical_errors) > len(related_critical_errors)
            else []
        ),
        "payload_catalog_validation": {
            "authorized": False,
            "preview_only": True,
            "reason": (
                "Catalog identity review does not prove official technician "
                "stock or group distribution"
            ),
        },
        "safety": {
            "network_used": False,
            "post_executed": False,
            "datasnap_write_executed": False,
            "stock_movement_executed": False,
        },
    }


def render_report_text(report: dict) -> str:
    lines = [
        "RELATORIO OFFLINE DE CATALOGO OFICIAL - CONTRATO 2221170",
        "",
        f"Arquivo: {report['catalog']['metadata']['source_file']}",
        f"SHA-256: {report['catalog']['metadata']['source_sha256']}",
        f"Registros: {report['catalog']['metadata']['record_count']}",
        "",
    ]
    for material in report["materials"]:
        official_description = material["official"]["description"] or "-"
        candidates = ", ".join(
            candidate["code"]
            for candidate in material[
                "official_candidates_by_exact_normalized_description"
            ]
        ) or "-"
        lines.extend(
            [
                f"{material['toa']['code']} | {material['status'].upper()}",
                f"  TOA: {material['toa']['description']}",
                (
                    "  Desktop: "
                    f"{material['desktop']['code']} "
                    f"{material['desktop']['description']}"
                ),
                f"  Oficial: {official_description}",
                f"  Candidatos por descricao exata: {candidates}",
                f"  Motivos: {', '.join(material['reasons']) or '-'}",
            ]
        )
    lines.extend(
        [
            "",
            "DESTAQUES",
            "- 22056408 = CONECTOR ATENUADOR 06DB.",
            (
                "- FITA ISOLANTE 3M HIGHLAND 19MM X 20M = 22064608, "
                "candidato por descricao; estoque oficial nao comprovado."
            ),
            (
                "- O catalogo nao contem saldo nem grupo e nao autoriza "
                "distribuicao ou envio."
            ),
            "",
            "BLOCKERS",
        ]
    )
    lines.extend(f"- {blocker}" for blocker in report["blockers"])
    return "\n".join(lines) + "\n"


def _assert_outside_logs(path: Path) -> None:
    if any(part.casefold() == "logs" for part in path.resolve().parts):
        raise CatalogValidationError(
            "Catalogo, indice e relatorio devem permanecer fora de logs"
        )


def _write_text_atomic(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(content, encoding="utf-8")
    temporary.replace(path)


def _write_json_atomic(path: Path, payload: dict) -> None:
    _write_text_atomic(
        path,
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
    )


def generate_offline_artifacts(
    *,
    xlsx_path: str | Path,
    index_output: str | Path,
    report_output: str | Path,
    text_output: str | Path | None = None,
) -> tuple[OfficialMaterialCatalog, dict]:
    source = Path(xlsx_path)
    index_path = Path(index_output)
    report_path = Path(report_output)
    _assert_outside_logs(source)
    _assert_outside_logs(index_path)
    _assert_outside_logs(report_path)
    if text_output is not None:
        _assert_outside_logs(Path(text_output))

    catalog = load_xlsx_catalog(source)
    report = build_contract_2221170_report(catalog)
    _write_json_atomic(index_path, catalog.to_index_dict())
    _write_json_atomic(report_path, report)
    if text_output is not None:
        _write_text_atomic(Path(text_output), render_report_text(report))
    return catalog, report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Importa e audita o catalogo oficial sem usar rede.",
    )
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--index-output", required=True)
    parser.add_argument("--report-output", required=True)
    parser.add_argument("--text-output")
    args = parser.parse_args(argv)

    catalog, report = generate_offline_artifacts(
        xlsx_path=args.xlsx,
        index_output=args.index_output,
        report_output=args.report_output,
        text_output=args.text_output,
    )
    print(
        json.dumps(
            {
                "metadata": catalog.metadata,
                "blockers": report["blockers"],
                "network_used": False,
                "authorized_for_send": False,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
